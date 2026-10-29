"""读取学生账号文件（支持 .xlsx 与 .csv）。

文件放在程序所在目录，文件名固定为 ``学生账号.xlsx`` 或 ``学生账号.csv``：

- 第一行可以是表头（如 ``学号,密码``），也可以直接写数据；
- 第一列是学号，第二列是密码；
- ``.xlsx`` 与 ``.csv`` 同时存在时优先使用 ``.xlsx``。
"""

from __future__ import annotations

import csv
import sys
from datetime import date, datetime
from pathlib import Path
from typing import Iterable

# 支持的文件名，按优先级排列
STUDENT_FILE_NAMES: tuple[str, ...] = ("学生账号.xlsx", "学生账号.csv")

# CSV 依次尝试的编码（UTF-8 含 BOM / GBK）
CSV_ENCODINGS: tuple[str, ...] = ("utf-8-sig", "gbk")

# 表头判断：中文关键字按包含匹配，英文按整格精确匹配
_HEADER_HINTS_CN: tuple[str, ...] = (
    "学号",
    "账号",
    "帐号",
    "用户名",
    "姓名",
    "登录名",
    "密码",
)
_HEADER_HINTS_EN: frozenset[str] = frozenset(
    {"name", "login", "username", "password", "pwd", "student", "id"}
)


def program_dir() -> Path:
    """程序所在目录：打包后为 exe 所在目录，否则为脚本所在目录。"""
    if getattr(sys, "frozen", False):
        return Path(sys.executable).resolve().parent
    return Path(__file__).resolve().parent


def search_dirs() -> list[Path]:
    """依次查找的目录：程序目录优先，其次当前工作目录。"""
    dirs: list[Path] = []
    for directory in (program_dir(), Path.cwd()):
        resolved = directory.resolve()
        if resolved not in dirs:
            dirs.append(resolved)
    return dirs


def find_student_file() -> Path:
    """按优先级查找学生账号文件，找不到时抛出 FileNotFoundError。"""
    dirs = search_dirs()
    for directory in dirs:
        for name in STUDENT_FILE_NAMES:
            candidate = directory / name
            if candidate.is_file():
                return candidate
    searched = "\n".join(f"  - {directory}" for directory in dirs)
    raise FileNotFoundError(
        "未找到学生账号文件，请将 "
        + " 或 ".join(STUDENT_FILE_NAMES)
        + " 放到程序所在目录。\n已查找目录：\n"
        + searched
    )


def _cell_text(value: object) -> str:
    """把单元格值转成字符串，避免 Excel 数字学号变成 ``1234.0``。"""
    if value is None:
        return ""
    if isinstance(value, bool):
        return str(value)
    if isinstance(value, float) and value.is_integer():
        return str(int(value))
    if isinstance(value, (datetime, date)):
        return value.strftime("%Y-%m-%d")
    return str(value)


def _cells_to_pair(cells: Iterable[object] | None) -> tuple[str, str]:
    """取一行的前两列，返回 (学号, 密码)，并去掉首尾空白。"""
    values = list(cells) if cells else []
    name = _cell_text(values[0]) if len(values) > 0 else ""
    password = _cell_text(values[1]) if len(values) > 1 else ""
    return name.strip(), password.strip()


def _read_xlsx_rows(path: Path) -> list[tuple[str, str]]:
    from openpyxl import load_workbook

    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        return [
            _cells_to_pair(row)
            for row in workbook.active.iter_rows(values_only=True)
        ]
    finally:
        workbook.close()


def _read_csv_rows(path: Path) -> list[tuple[str, str]]:
    last_error: UnicodeDecodeError | None = None
    for encoding in CSV_ENCODINGS:
        try:
            with path.open("r", encoding=encoding, newline="") as file:
                return [_cells_to_pair(row) for row in csv.reader(file)]
        except UnicodeDecodeError as exc:
            last_error = exc
    raise ValueError(
        f"无法识别 {path.name} 的编码，请另存为 UTF-8 或 ANSI(GBK) 后重试。"
    ) from last_error


def _is_header_cell(text: str) -> bool:
    """判断单元格是否为表头字样。"""
    if not text:
        return False
    if text.lower() in _HEADER_HINTS_EN:
        return True
    return any(hint in text for hint in _HEADER_HINTS_CN)


def _is_header_row(name: str, password: str) -> bool:
    """判断一行是否为表头。

    必须是"学号列是表头字样"，且密码列为空或也是表头字样——
    密码里出现"密码"等字样（例如"中文密码"）不能被当成表头。
    """
    if not _is_header_cell(name):
        return False
    return not password or _is_header_cell(password)


def load_students() -> dict[str, str]:
    """读取学生账号文件，返回 ``{学号: 密码}``。

    找不到文件时抛出 :class:`FileNotFoundError`；格式错误时抛出其他异常。
    """
    path = find_student_file()
    try:
        if path.suffix.lower() == ".xlsx":
            rows = _read_xlsx_rows(path)
        else:
            rows = _read_csv_rows(path)
    except Exception as exc:
        raise RuntimeError(
            f"无法读取 {path.name}（{exc}）。"
            "请确认文件未损坏，且格式为 Excel 2007 及以上的 .xlsx 或 .csv。"
        ) from exc

    students: dict[str, str] = {}
    incomplete = 0
    for name, password in rows:
        # 表头行任意位置都可跳过（Excel 里常见标题行、重复表头）
        if _is_header_row(name, password):
            continue
        if not name or not password:
            if name or password:
                incomplete += 1
            continue
        students[name] = password

    if incomplete:
        print(f"警告：{path.name} 中有 {incomplete} 行缺少学号或密码，已跳过。")
    print(f"已读取 {path.name}，共 {len(students)} 名学生。")
    return students
